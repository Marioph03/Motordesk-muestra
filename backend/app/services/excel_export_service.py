"""Generacion de libros Excel descargables para los informes economicos."""

from datetime import timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class ExcelExportService:
    """Convierte los datos ya calculados de informes en un XLSX legible."""

    _AZUL = "123A63"
    _AZUL_CLARO = "E8F0F8"

    @staticmethod
    def generar_listado(titulo, cabecera, filas, anchos):
        """Genera una exportaciÃ³n reutilizable de datos maestros."""
        libro = Workbook()
        hoja = libro.active
        hoja.title = titulo[:31]
        for columna, texto in enumerate(cabecera, start=1):
            celda = hoja.cell(1, columna, texto)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=ExcelExportService._AZUL)
            celda.alignment = Alignment(horizontal="center")
        for fila_indice, valores in enumerate(filas, start=2):
            for columna, valor in enumerate(valores, start=1):
                celda = hoja.cell(fila_indice, columna, valor)
                if isinstance(valor, (int, float)) and columna > 1:
                    celda.alignment = Alignment(horizontal="right")
        for columna, ancho in enumerate(anchos, start=1):
            hoja.column_dimensions[get_column_letter(columna)].width = ancho
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = f"A1:{get_column_letter(len(cabecera))}{max(1, len(filas) + 1)}"
        salida = BytesIO()
        libro.save(salida)
        salida.seek(0)
        return salida

    @staticmethod
    def generar_informe_economico(titulo_informe, informe, empresa):
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Informe"

        es_clientes = titulo_informe == "RESUMEN ECONOMICO POR CLIENTE"
        es_compras = titulo_informe == "COMPRAS Y GASTOS"
        es_facturacion = not es_clientes and not es_compras

        if es_clientes:
            cabecera = ["Cliente", "NIF/CIF", "Facturas", "Base", "IVA", "Total"]
            anchos = [30, 16, 12, 14, 14, 14]
        elif es_compras:
            cabecera = ["Fecha", "Factura", "Tipo", "Proveedor", "NIF/CIF", "Base", "IVA", "Total"]
            anchos = [14, 18, 16, 28, 16, 14, 14, 14]
        else:
            cabecera = ["Fecha", "Factura", "Cliente", "NIF/CIF", "Base", "IVA", "Total", "Metodo de pago"]
            anchos = [14, 18, 28, 16, 14, 14, 14, 22]

        ultima_columna = len(cabecera)
        fecha_fin = informe["fecha_fin"] - timedelta(days=1)
        nombre_empresa = getattr(empresa, "nombre", None) or getattr(empresa, "nombre_comercial", None) or "MotorDesk"

        hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_columna)
        titulo = hoja.cell(1, 1, titulo_informe)
        titulo.font = Font(bold=True, size=16, color="FFFFFF")
        titulo.fill = PatternFill("solid", fgColor=ExcelExportService._AZUL)
        titulo.alignment = Alignment(horizontal="center")

        hoja.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima_columna)
        periodo = hoja.cell(2, 1, f"{nombre_empresa} - Periodo: {informe['fecha_inicio']:%d/%m/%Y} - {fecha_fin:%d/%m/%Y}")
        periodo.font = Font(italic=True, color="5B6470")

        fila_cabecera = 4
        for columna, texto in enumerate(cabecera, start=1):
            celda = hoja.cell(fila_cabecera, columna, texto)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=ExcelExportService._AZUL)
            celda.alignment = Alignment(horizontal="center")

        for indice, fila in enumerate(informe["filas"], start=fila_cabecera + 1):
            if es_clientes:
                valores = [fila["tercero"], fila["nif"], fila["documentos"], fila["base"], fila["iva"], fila["total"]]
            elif es_compras:
                valores = [fila["fecha"], fila["numero"], fila.get("tipo_gasto", "COMPRA").title(), fila["tercero"], fila["nif"], fila["base"], fila["iva"], fila["total"]]
            else:
                valores = [fila["fecha"], fila["numero"], fila["tercero"], fila["nif"], fila["base"], fila["iva"], fila["total"], fila.get("metodo_pago", "Pendiente")]

            for columna, valor in enumerate(valores, start=1):
                celda = hoja.cell(indice, columna, valor)
                if columna in ([4, 5, 6] if es_clientes else ([6, 7, 8] if es_compras else [5, 6, 7])):
                    celda.number_format = '#,##0.00 "EUR"'
                    celda.alignment = Alignment(horizontal="right")
                elif columna == 1 and not es_clientes:
                    celda.number_format = "DD/MM/YYYY"

        fila_totales = max(fila_cabecera + 2, fila_cabecera + len(informe["filas"]) + 2)
        hoja.cell(fila_totales, ultima_columna - 1, "Documentos" if not es_clientes else "Clientes").font = Font(bold=True)
        hoja.cell(fila_totales, ultima_columna, informe["totales"]["documentos"])
        for desplazamiento, etiqueta, clave in [(1, "Base imponible", "base"), (2, "IVA", "iva"), (3, "TOTAL", "total")]:
            hoja.cell(fila_totales + desplazamiento, ultima_columna - 1, etiqueta).font = Font(bold=etiqueta == "TOTAL")
            valor = hoja.cell(fila_totales + desplazamiento, ultima_columna, informe["totales"][clave])
            valor.font = Font(bold=etiqueta == "TOTAL")
            valor.number_format = '#,##0.00 "EUR"'
            if etiqueta == "TOTAL":
                for columna in range(ultima_columna - 1, ultima_columna + 1):
                    hoja.cell(fila_totales + desplazamiento, columna).fill = PatternFill("solid", fgColor=ExcelExportService._AZUL_CLARO)

        for columna, ancho in enumerate(anchos, start=1):
            hoja.column_dimensions[get_column_letter(columna)].width = ancho
        hoja.freeze_panes = "A5"
        hoja.auto_filter.ref = f"A{fila_cabecera}:{get_column_letter(ultima_columna)}{max(fila_cabecera, fila_cabecera + len(informe['filas']))}"

        salida = BytesIO()
        libro.save(salida)
        salida.seek(0)
        return salida
