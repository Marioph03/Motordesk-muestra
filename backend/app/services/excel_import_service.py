"""Análisis e importación controlada de hojas Excel heredadas hacia MotorDesk."""

from pathlib import Path
from uuid import uuid4

from fastapi import UploadFile

from openpyxl import (
    load_workbook
)
from sqlalchemy.orm import Session

from app.models.articulo_stock import ArticuloStock
from app.models.cliente import Cliente
from app.models.proveedor import Proveedor
from app.models.vehiculo import Vehiculo


class ExcelImportService:
    """Mantiene archivos temporales y valida el mapeo antes de insertar clientes."""


    TEMP_DIR = Path(
        "storage/importaciones/temp"
    )

    EXTENSIONES_VALIDAS = {
        ".xlsx"
    }


    @staticmethod
    async def guardar_temporal(
        archivo: UploadFile
    ) -> Path:

        nombre_original = (
            archivo.filename
            or ""
        )

        extension = (
            Path(nombre_original)
            .suffix
            .lower()
        )


        if (
            extension
            not in
            ExcelImportService
            .EXTENSIONES_VALIDAS
        ):

            raise ValueError(
                "Solo se permiten archivos XLSX"
            )


        contenido = await archivo.read()


        # Máximo 25 MB
        if (
            len(contenido)
            >
            25 * 1024 * 1024
        ):

            raise ValueError(
                "El Excel no puede superar los 25 MB"
            )


        ExcelImportService.TEMP_DIR.mkdir(
            parents=True,
            exist_ok=True
        )


        nombre = (
            f"{uuid4().hex}.xlsx"
        )


        ruta = (
            ExcelImportService
            .TEMP_DIR
            /
            nombre
        )


        ruta.write_bytes(
            contenido
        )


        return ruta


    @staticmethod
    def analizar(
        ruta: Path,
        nombre_hoja: str | None = None,
    ) -> dict:

        if not ruta.exists():

            raise ValueError(
                "El archivo temporal no existe"
            )


        try:
            workbook = load_workbook(filename=ruta, read_only=True, data_only=True)
        except Exception as error:
            raise ValueError("El archivo no es un Excel XLSX vÃ¡lido") from error


        try:
            if not workbook.sheetnames:
                raise ValueError("El archivo no contiene hojas")

            hojas = []


            for nombre in workbook.sheetnames:

                hoja = workbook[nombre]

                hojas.append({

                    "nombre": nombre,

                "filas":
                    hoja.max_row,

                "columnas":
                    hoja.max_column

                })


            if nombre_hoja and nombre_hoja not in workbook.sheetnames:
                raise ValueError("La hoja seleccionada no existe")
            hoja_activa = workbook[nombre_hoja or workbook.sheetnames[0]]


            cabecera = []
            primera_fila = next(hoja_activa.iter_rows(min_row=1, max_row=1), None)
            if primera_fila is None:
                raise ValueError("La hoja seleccionada está vacía")


            for indice, celda in enumerate(
            primera_fila,
            start=1
            ):

                valor = (
                str(celda.value).strip()
                if celda.value
                is not None
                else f"Columna {indice}"
                )


                cabecera.append(
                {
                    "indice":
                        indice,

                    "nombre":
                        valor
                }
                )


            nombres_columnas = [
            item["nombre"]
            for item in cabecera
            ]


            muestra = []


            for fila in hoja_activa.iter_rows(
            min_row=2,
            max_row=min(
                hoja_activa.max_row,
                11
            ),
            values_only=True
            ):

                registro = {}


                for indice, valor in enumerate(
                fila
                ):

                    if (
                    indice
                    <
                    len(
                        nombres_columnas
                    )
                    ):

                        registro[
                        nombres_columnas[
                            indice
                        ]
                        ] = (
                        ExcelImportService
                        .serializar_valor(
                            valor
                        )
                        )


                muestra.append(
                registro
                )


            return {

            "archivo":
                ruta.name,

            "hojas":
                hojas,

            "hoja_activa":
                hoja_activa.title,

            "columnas":
                cabecera,

            "total_filas":
                max(
                    hoja_activa.max_row - 1,
                    0
                ),

            "muestra":
                muestra

            }
        finally:
            workbook.close()


    @staticmethod
    def serializar_valor(
        valor
    ):

        if valor is None:
            return None


        if hasattr(
            valor,
            "isoformat"
        ):

            return (
                valor.isoformat()
            )


        if isinstance(
            valor,
            (
                str,
                int,
                float,
                bool
            )
        ):

            return valor


        return str(
            valor
        )


    @staticmethod
    def obtener_ruta_temporal(
        nombre: str
    ) -> Path:

        # Evitamos que alguien pueda
        # enviar ../../archivo.xlsx

        nombre_seguro = (
            Path(nombre).name
        )


        ruta = (
            ExcelImportService
            .TEMP_DIR
            /
            nombre_seguro
        )


        if not ruta.exists():

            raise ValueError(
                "Archivo de importación "
                "no encontrado"
            )


        return ruta

    @staticmethod
    def normalizar_texto(valor) -> str | None:
        if valor is None:
            return None
        texto = str(valor).strip()
        return texto or None

    @staticmethod
    def normalizar_documento(valor) -> str | None:
        texto = ExcelImportService.normalizar_texto(valor)
        return texto.upper().replace(" ", "").replace("-", "").replace(".", "") if texto else None

    @staticmethod
    def normalizar_telefono(valor) -> str | None:
        texto = ExcelImportService.normalizar_texto(valor)
        return texto.replace(" ", "").replace("-", "").replace(".", "") if texto else None

    @staticmethod
    def normalizar_email(valor) -> str | None:
        texto = ExcelImportService.normalizar_texto(valor)
        return texto.lower() if texto else None

    @staticmethod
    def leer_registros_mapeados(ruta: Path, hoja_nombre: str, mapeo: list) -> list[dict]:
        workbook = load_workbook(filename=ruta, read_only=True, data_only=True)
        try:
            if hoja_nombre not in workbook.sheetnames:
                raise ValueError("La hoja seleccionada no existe")
            hoja = workbook[hoja_nombre]
            cabecera = next(hoja.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if cabecera is None:
                return []
            columnas = {str(valor).strip(): indice for indice, valor in enumerate(cabecera) if valor is not None}
            mapa_campos = {}
            for item in mapeo:
                campo = getattr(item, "campo_motordesk", None)
                columna = getattr(item, "columna_excel", None)
                if not campo:
                    continue
                if columna and columna not in columnas:
                    raise ValueError(f"La columna '{columna}' no existe en el Excel")
                mapa_campos[campo] = columnas[columna] if columna else None

            registros = []
            for numero_fila, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
                datos = {}
                fila_vacia = True
                for campo, indice in mapa_campos.items():
                    valor = fila[indice] if indice is not None and indice < len(fila) else None
                    datos[campo] = valor
                    if valor is not None and str(valor).strip():
                        fila_vacia = False
                if not fila_vacia:
                    registros.append({"numero_fila": numero_fila, "datos": datos})
            return registros
        finally:
            workbook.close()

    @staticmethod
    def validar_cliente_preview(db: Session, numero_fila: int, datos: dict) -> dict:
        errores: list[str] = []
        advertencias: list[str] = []
        nombre = ExcelImportService.normalizar_texto(datos.get("nombre"))
        dni_nif = ExcelImportService.normalizar_documento(datos.get("dni_nif"))
        telefono = ExcelImportService.normalizar_telefono(datos.get("telefono"))
        email = ExcelImportService.normalizar_email(datos.get("email"))
        if not nombre:
            errores.append("El nombre es obligatorio")
        if email and ("@" not in email or "." not in email.split("@")[-1]):
            advertencias.append("El email parece tener un formato incorrecto")
        if telefono and len(telefono) < 6:
            advertencias.append("El teléfono parece demasiado corto")
        duplicado = bool(dni_nif and db.query(Cliente).filter(Cliente.dni == dni_nif).first())
        if duplicado:
            advertencias.append("Ya existe un cliente con este DNI/NIF")
        datos_limpios = {
            "nombre": nombre,
            "apellidos": ExcelImportService.normalizar_texto(datos.get("apellidos")),
            "dni_nif": dni_nif,
            "telefono": telefono,
            "email": email,
            "direccion": ExcelImportService.normalizar_texto(datos.get("direccion")),
            "codigo_postal": ExcelImportService.normalizar_texto(datos.get("codigo_postal")),
            "localidad": ExcelImportService.normalizar_texto(datos.get("localidad")),
            "provincia": ExcelImportService.normalizar_texto(datos.get("provincia")),
            "observaciones": ExcelImportService.normalizar_texto(datos.get("observaciones")),
        }
        return {"numero_fila": numero_fila, "datos": datos_limpios, "valida": not errores and not duplicado, "duplicado": duplicado, "errores": errores, "advertencias": advertencias}

    @staticmethod
    def preview_clientes(db: Session, archivo_temporal: str, hoja: str, mapeo: list) -> dict:
        registros = ExcelImportService.leer_registros_mapeados(ExcelImportService.obtener_ruta_temporal(archivo_temporal), hoja, mapeo)
        filas = [ExcelImportService.validar_cliente_preview(db, registro["numero_fila"], registro["datos"]) for registro in registros]
        return {
            "total": len(filas),
            "validas": sum(1 for fila in filas if fila["valida"]),
            "errores": sum(1 for fila in filas if fila["errores"]),
            "advertencias": sum(1 for fila in filas if fila["advertencias"] and not fila["errores"]),
            "duplicados": sum(1 for fila in filas if fila["duplicado"]),
            "filas": filas,
        }

    @staticmethod
    def importar_clientes(
        db: Session,
        archivo_temporal: str,
        hoja: str,
        mapeo: list,
        ignorar_duplicados: bool = True,
    ) -> dict:
        registros = ExcelImportService.leer_registros_mapeados(
            ExcelImportService.obtener_ruta_temporal(archivo_temporal),
            hoja,
            mapeo,
        )
        importados = omitidos = errores = 0
        detalle_errores: list[str] = []

        try:
            for registro in registros:
                numero_fila = registro["numero_fila"]
                resultado = ExcelImportService.validar_cliente_preview(
                    db,
                    numero_fila,
                    registro["datos"],
                )
                if resultado["errores"]:
                    errores += 1
                    detalle_errores.append(
                        f"Fila {numero_fila}: {'; '.join(resultado['errores'])}"
                    )
                    continue
                if resultado["duplicado"]:
                    if ignorar_duplicados:
                        omitidos += 1
                    else:
                        errores += 1
                        detalle_errores.append(f"Fila {numero_fila}: cliente duplicado")
                    continue

                datos = resultado["datos"]
                try:
                    with db.begin_nested():
                        # El modelo actual usa dni y ciudad, no dni_nif/localidad.
                        cliente = Cliente(
                            nombre=datos["nombre"],
                            apellidos=datos["apellidos"],
                            dni=datos["dni_nif"],
                            telefono=datos["telefono"],
                            email=datos["email"],
                            direccion=datos["direccion"],
                            codigo_postal=datos["codigo_postal"],
                            ciudad=datos["localidad"],
                            observaciones=datos["observaciones"],
                            activo=True,
                        )
                        db.add(cliente)
                        db.flush()
                    importados += 1
                except Exception as error:
                    errores += 1
                    detalle_errores.append(f"Fila {numero_fila}: {str(error)}")
            db.commit()
        except Exception:
            db.rollback()
            raise

        return {
            "total": len(registros),
            "importados": importados,
            "omitidos": omitidos,
            "errores": errores,
            "detalle_errores": detalle_errores,
        }

    @staticmethod
    def preview(db: Session, tipo: str, archivo_temporal: str, hoja: str, mapeo: list) -> dict:
        registros = ExcelImportService.leer_registros_mapeados(
            ExcelImportService.obtener_ruta_temporal(archivo_temporal), hoja, mapeo
        )
        filas = [ExcelImportService._validar_registro(db, tipo, item["numero_fila"], item["datos"]) for item in registros]
        return {
            "total": len(filas),
            "validas": sum(1 for fila in filas if fila["valida"]),
            "errores": sum(1 for fila in filas if fila["errores"]),
            "advertencias": sum(1 for fila in filas if fila["advertencias"] and not fila["errores"]),
            "duplicados": sum(1 for fila in filas if fila["duplicado"]),
            "filas": filas,
        }

    @staticmethod
    def importar(db: Session, tipo: str, archivo_temporal: str, hoja: str, mapeo: list, ignorar_duplicados: bool = True) -> dict:
        registros = ExcelImportService.leer_registros_mapeados(
            ExcelImportService.obtener_ruta_temporal(archivo_temporal), hoja, mapeo
        )
        importados = omitidos = errores = 0
        detalle_errores: list[str] = []
        try:
            for registro in registros:
                resultado = ExcelImportService._validar_registro(db, tipo, registro["numero_fila"], registro["datos"])
                if resultado["errores"]:
                    errores += 1
                    detalle_errores.append(f"Fila {registro['numero_fila']}: {'; '.join(resultado['errores'])}")
                    continue
                if resultado["duplicado"]:
                    if ignorar_duplicados:
                        omitidos += 1
                    else:
                        errores += 1
                        detalle_errores.append(f"Fila {registro['numero_fila']}: registro duplicado")
                    continue
                try:
                    with db.begin_nested():
                        ExcelImportService._crear_registro(db, tipo, resultado["datos"])
                    importados += 1
                except Exception as error:
                    errores += 1
                    detalle_errores.append(f"Fila {registro['numero_fila']}: {str(error)}")
            db.commit()
        except Exception:
            db.rollback()
            raise
        return {"total": len(registros), "importados": importados, "omitidos": omitidos, "errores": errores, "detalle_errores": detalle_errores}

    @staticmethod
    def _validar_registro(db: Session, tipo: str, numero_fila: int, datos: dict) -> dict:
        if tipo == "clientes":
            return ExcelImportService.validar_cliente_preview(db, numero_fila, datos)
        errores: list[str] = []
        advertencias: list[str] = []
        duplicado = False
        limpio = {campo: ExcelImportService.normalizar_texto(valor) for campo, valor in datos.items()}
        if tipo == "vehiculos":
            limpio["matricula"] = (limpio.get("matricula") or "").upper().replace(" ", "")
            limpio["cliente_dni_nif"] = ExcelImportService.normalizar_documento(datos.get("cliente_dni_nif"))
            if not limpio["matricula"]:
                errores.append("La matricula es obligatoria")
            if not limpio["cliente_dni_nif"]:
                errores.append("El DNI/NIF del propietario es obligatorio")
            elif db.query(Cliente).filter(Cliente.dni == limpio["cliente_dni_nif"]).first() is None:
                errores.append("No existe un cliente con ese DNI/NIF")
            duplicado = bool(limpio["matricula"] and db.query(Vehiculo).filter(Vehiculo.matricula == limpio["matricula"]).first())
        elif tipo == "proveedores":
            limpio["nif_cif"] = ExcelImportService.normalizar_documento(datos.get("nif_cif"))
            if not limpio.get("nombre"):
                errores.append("El nombre es obligatorio")
            if not limpio["nif_cif"]:
                errores.append("El NIF/CIF es obligatorio")
            duplicado = bool(limpio["nif_cif"] and db.query(Proveedor).filter(Proveedor.cif == limpio["nif_cif"]).first())
        elif tipo == "articulos":
            limpio["referencia"] = (limpio.get("referencia") or "").upper()
            if not limpio["referencia"]:
                errores.append("La referencia es obligatoria")
            if not limpio.get("nombre"):
                errores.append("El nombre es obligatorio")
            duplicado = bool(limpio["referencia"] and db.query(ArticuloStock).filter(ArticuloStock.referencia == limpio["referencia"]).first())
        else:
            raise ValueError("Tipo de importacion no valido")
        if duplicado:
            advertencias.append("Ya existe un registro con ese identificador")
        return {"numero_fila": numero_fila, "datos": limpio, "valida": not errores and not duplicado, "duplicado": duplicado, "errores": errores, "advertencias": advertencias}

    @staticmethod
    def _numero(valor) -> float:
        if valor is None or str(valor).strip() == "":
            return 0.0
        return float(str(valor).replace(",", "."))

    @staticmethod
    def _crear_registro(db: Session, tipo: str, datos: dict) -> None:
        if tipo == "clientes":
            cliente = Cliente(nombre=datos["nombre"], apellidos=datos.get("apellidos"), dni=datos.get("dni_nif"), telefono=datos.get("telefono"), email=datos.get("email"), direccion=datos.get("direccion"), codigo_postal=datos.get("codigo_postal"), ciudad=datos.get("localidad"), provincia=datos.get("provincia"), observaciones=datos.get("observaciones"), activo=True)
            db.add(cliente); db.flush(); cliente.codigo = f"CL-{cliente.id:06d}"
        elif tipo == "vehiculos":
            cliente = db.query(Cliente).filter(Cliente.dni == datos["cliente_dni_nif"]).first()
            vehiculo = Vehiculo(cliente_id=cliente.id, matricula=datos["matricula"], marca=datos.get("marca") or "Sin especificar", modelo=datos.get("modelo") or "Sin especificar", version=datos.get("version"), bastidor=datos.get("bastidor"), kilometros=int(ExcelImportService._numero(datos.get("kilometros"))), combustible=datos.get("combustible"))
            db.add(vehiculo); db.flush(); vehiculo.codigo = f"VEH-{vehiculo.id:06d}"
        elif tipo == "proveedores":
            proveedor = Proveedor(nombre=datos["nombre"], cif=datos["nif_cif"], telefono=datos.get("telefono"), email=datos.get("email"), direccion=datos.get("direccion"), ciudad=datos.get("localidad"), provincia=datos.get("provincia"), activo=True)
            db.add(proveedor); db.flush(); proveedor.codigo = f"PRO-{proveedor.id:06d}"
        elif tipo == "articulos":
            db.add(ArticuloStock(referencia=datos["referencia"], nombre=datos["nombre"], descripcion=datos.get("descripcion"), stock_actual=ExcelImportService._numero(datos.get("stock_actual")), stock_minimo=ExcelImportService._numero(datos.get("stock_minimo")), coste=ExcelImportService._numero(datos.get("coste_unitario")), precio_venta=ExcelImportService._numero(datos.get("coste_unitario"))))
